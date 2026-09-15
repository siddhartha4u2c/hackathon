import os
import sys
from pathlib import Path

import openpyxl
import psycopg

TABLES = {
    "Dealers": ("dealers", ["dealer_id", "dealer_name", "region", "country", "city", "tier", "status", "contact_email", "credit_limit_eur", "onboarded_date"]),
    "Parts": ("parts", ["part_no", "description", "category", "unit_price_eur", "warranty_months", "supplier", "status"]),
    "PurchaseOrders": ("purchase_orders", ["po_no", "po_line_no", "dealer_id", "part_no", "order_qty", "unit_price_eur", "line_total_eur", "order_date", "requested_delivery_date", "po_status", "shipment_id", "currency"]),
    "Shipments": ("shipments", ["shipment_id", "po_no", "dealer_id", "carrier", "ship_date", "est_delivery_date", "actual_delivery_date", "shipment_status", "tracking_no", "qty"]),
    "Claims": ("claims", ["claim_id", "claim_type", "dealer_id", "part_no", "po_no", "claim_qty", "purchase_date", "claim_date", "claim_status", "reason", "claim_amount_eur", "resolution_code"]),
    "Inventory": ("inventory", ["inventory_id", "part_no", "warehouse_loc", "on_hand_qty", "reserved_qty", "available_qty", "reorder_point", "bin_location", "last_count_date"]),
    "BOM": ("bom", ["bom_id", "assembly_part_no", "component_part_no", "qty_per", "level", "notes"]),
    "Knowledge": ("knowledge", ["doc_id", "doc_type", "title", "module", "error_code", "summary", "last_updated", "owner_team"]),
}

SOURCE_ALIASES = {
    "description": ("description", "partname"),
    "supplier": ("supplier", "suppliername", "supplierid"),
}


def clean(value):
    return None if value in ("", None) else value


def main():
    workbook_path = Path(sys.argv[1] if len(sys.argv) > 1 else "/app/data/dataset.xlsx")
    if not workbook_path.exists():
        raise FileNotFoundError(workbook_path)
    with psycopg.connect(os.environ["DATABASE_URL"]) as conn:
        with conn.cursor() as cur:
            schema = Path(__file__).parent.parent / "app" / "schema.sql"
            cur.execute(schema.read_text())
            for sheet, (table, columns) in TABLES.items():
                cur.execute(f"TRUNCATE {table} CASCADE")
                ws = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)[sheet]
                headers = [str(cell.value).strip() for cell in next(ws.iter_rows())]
                normalized = {h.lower().replace(" ", "").replace("_", ""): i for i, h in enumerate(headers)}
                indexes = []
                for column in columns:
                    candidates = SOURCE_ALIASES.get(column, (column,))
                    candidates = tuple(candidate.lower().replace(" ", "").replace("_", "") for candidate in candidates)
                    source_index = next(
                        (normalized[candidate] for candidate in candidates if candidate in normalized),
                        None,
                    )
                    if source_index is None:
                        raise ValueError(f"Sheet {sheet} is missing a source column for {column}")
                    indexes.append(source_index)
                quoted = ", ".join(columns)
                placeholders = ", ".join(["%s"] * len(columns))
                for row in ws.iter_rows(min_row=2, values_only=True):
                    values = [clean(row[index]) for index in indexes]
                    if any(value is not None for value in values):
                        cur.execute(f"INSERT INTO {table} ({quoted}) VALUES ({placeholders}) ON CONFLICT DO NOTHING", values)
        conn.commit()
    print(f"Imported {workbook_path}")


if __name__ == "__main__":
    main()
